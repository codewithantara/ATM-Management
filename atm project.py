from openpyxl import load_workbook
"""
wb=load_workbook("atm_data.xlsx")
ws=wb.active
print("excel active")
print("A1:",ws["A1"].value)
print("B1:",ws["B1"].value)
wb.close()
"""

"""
#add new account
def add_new():
    mobile_no=int(input("enter mobile number:"))
    amount=float(input("enter amount:"))
    wb=load_workbook("atm_data.xlsx")
    ws=wb.active
    ws.append([mobile_no,amount])
    wb.save("atm_data.xlsx")
    wb.close()
    print("Account created successfully")
add_new()
"""

"""
#deposit money
def deposit():
    mobile_no=int(input("enter mobile no:"))
    deposit_amount=float(input("enter amount:"))
    wb=load_workbook("atm_data.xlsx")
    ws=wb.active
    for i in range(2,ws.max_row+1):
        if str(ws.cell(i, 1).value).strip() == str(mobile_no).strip():
            balance=ws.cell(i,2).value
            ws.cell(i,2).value=balance+deposit_amount
            wb.save("atm_data.xlsx")
            wb.close()
            print("Money deposited successfully")
            return
        wb.close()
        print("Account not found")
deposit()
"""

"""
# withdraw money
def withdraw():
    mobile_no = input("enter mobile no:")
    amount = float(input("enter amount:"))
    wb = load_workbook("atm_data.xlsx")
    ws = wb.active
    for i in range(2, ws.max_row + 1):
        if str(ws.cell(i, 1).value).strip() == mobile_no.strip():
            balance = ws.cell(i, 2).value
            if amount <= balance:
                ws.cell(i, 2).value = balance - amount
                wb.save("atm_data.xlsx")
                wb.close()
                print("Money withdrawn successfully!")
            else:
                wb.close()
                print("Insufficient balance!")
            return
    wb.close()
    print("Account not found")
withdraw()
"""



"""
# View balance
def view_balance():
    mobile_no = input("Enter mobile number: ")
    wb = load_workbook("atm_data.xlsx")
    ws = wb.active
    for i in range(2, ws.max_row + 1):
        if str(ws.cell(i, 1).value).strip() == mobile_no.strip():
            balance = ws.cell(i, 2).value
            print("Your current balance is:", balance)
            wb.close()
            return
    wb.close()
    print("Account not found!")
view_balance()
"""

"""
#exit
def exit_atm():
    print("Thank you for using ATM!")
    print("Exiting program...")
exit_atm()
"""








            

    

    
        
